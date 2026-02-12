import io
from typing import List


def parse_file_to_text(path: str, ext: str) -> str:
    ext = ext.lower()
    if ext in ("txt", "md"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == "pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception as e:
            raise RuntimeError("pypdf is required to parse pdf") from e
        reader = PdfReader(path)
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)

    if ext == "docx":
        try:
            import docx  # type: ignore
        except Exception as e:
            raise RuntimeError("python-docx is required to parse docx") from e
        doc = docx.Document(path)
        return "\n".join([p.text for p in doc.paragraphs])

    if ext == "xlsx":
        try:
            import openpyxl  # type: ignore
        except Exception as e:
            raise RuntimeError("openpyxl is required to parse xlsx") from e
        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = "\t".join([str(x) for x in row if x is not None])
                if line:
                    parts.append(line)
        return "\n".join(parts)

    raise ValueError("unsupported file type")


def chunk_text(text: str, chunk_size: int, overlap: int) -> List[str]:
    if chunk_size <= 0:
        return []
    text = text.replace("\r\n", "\n").strip()
    if not text:
        return []
    step = max(1, chunk_size - max(0, overlap))
    chunks = []
    i = 0
    n = len(text)
    while i < n:
        chunk = text[i : i + chunk_size]
        if chunk:
            chunks.append(chunk)
        i += step
    return chunks
